from pathlib import Path

import pymupdf
from docx import Document


def extract_pdf(
    path: Path,
) -> str:
    document = pymupdf.open(str(path))

    text_parts = []

    try:
        for page_number, page in enumerate(
            document,
            start=1,
        ):
            text = page.get_text("text").strip()

            if text:
                print(
                    f"  Page {page_number}/{len(document)}: "
                    f"PyMuPDF"
                )

                text_parts.append(text)

    finally:
        document.close()

    return "\n".join(text_parts)


def extract_docx(
    path: Path,
) -> str:
    document = Document(path)

    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                text_parts.append(
                    " | ".join(cells)
                )

    return "\n".join(text_parts)


def extract_text(
    path: Path,
) -> str:
    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )