from pathlib import Path

import pymupdf
from docx import Document


def extract_pdf(path: Path) -> str:
    document = pymupdf.open(str(path))

    text_parts: list[str] = []

    try:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()

            if text:
                print(
                    f"Page {page_number}/{len(document)}: "
                    "text extracted"
                )

                text_parts.append(text)

    finally:
        document.close()

    return "\n\n".join(text_parts)


def extract_docx(path: Path) -> str:
    document = Document(path)

    text_parts: list[str] = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                text_parts.append(" | ".join(cells))

    return "\n\n".join(text_parts)


def extract_text(path: Path) -> str:
    path = Path(path)

    if not path.exists():
        raise FileNotFoundError(
            f"Resume file not found: {path}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    raise ValueError(
        f"Unsupported resume format: {extension}. "
        "Supported formats: .pdf, .docx"
    )